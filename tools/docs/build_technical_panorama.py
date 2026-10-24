"""Build the current technical panorama as a Word document.

Use the Codex bundled Python (python-docx and Pillow), or an explicitly prepared
documentation environment. No Godot runtime files are modified.
"""
import argparse
import hashlib
import json
from pathlib import Path
import re
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from technical_panorama_content import SECTIONS

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'docs/Sandcastle项目技术全景.docx'
INDEX = ROOT / 'docs/technical-panorama.sources.json'
WORK = ROOT / 'test_output/technical_panorama_20261021'
BASE = 'scripts/simulation/unified_edit/'
BLACK = RGBColor(0, 0, 0)

def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def xml(tag, **attrs):
    item = OxmlElement('w:' + tag)
    for key, value in attrs.items(): item.set(qn('w:' + key), str(value))
    return item

def font(run, size=None, bold=None):
    run.font.name = 'Calibri'
    run.font.color.rgb = BLACK
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')

def breakable(text):
    return str(text).replace('res://', '').replace('/', '/\u200b').replace('_', '_\u200b')

def paragraph(doc, text, style=None):
    p = doc.add_paragraph(style=style)
    font(p.add_run(text))
    return p

def heading(doc, text, level, bookmark=None):
    p = doc.add_paragraph(style=f'Heading {level}')
    font(p.add_run(text), bold=True)
    if bookmark:
        p._p.insert(0, xml('bookmarkStart', id=bookmark[0], name=bookmark[1]))
        p._p.append(xml('bookmarkEnd', id=bookmark[0]))
    return p

def table(doc, headers, rows, widths=None):
    widths = widths or [170 / len(headers)] * len(headers)
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Mm(width)
    pr = t._tbl.tblPr
    borders = xml('tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        borders.append(xml(edge, val='single', sz=4, color='D9D9D9'))
    pr.append(borders)
    margins = xml('tblCellMar')
    for side in ['top', 'bottom']: margins.append(xml(side, w=85, type='dxa'))
    for side in ['left', 'right']: margins.append(xml(side, w=100, type='dxa'))
    pr.append(margins)
    for ri, data in enumerate([headers] + rows):
        row = t.rows[0] if ri == 0 else t.add_row()
        row._tr.get_or_add_trPr().append(xml('cantSplit'))
        if ri == 0: row._tr.get_or_add_trPr().append(xml('tblHeader'))
        for ci, (cell, text) in enumerate(zip(row.cells, data)):
            cell.width = Mm(widths[ci]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell._tc.get_or_add_tcPr().append(xml('shd', fill='DDE8F0' if ri == 0 else ('F7F9FB' if ri % 2 == 0 else 'FFFFFF')))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.12
            p.paragraph_format.keep_with_next = ri == 0
            p.alignment = (WD_ALIGN_PARAGRAPH.CENTER
                           if re.fullmatch(r'(?:-?\d+(?:\.\d+)?|S\d{3}|true|false)', str(text), re.I)
                           else WD_ALIGN_PARAGRAPH.LEFT)
            font(p.add_run(breakable(text)), 9.2, ri == 0)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.line_spacing = 0.35
    return t

def diagram(name):
    path = WORK / (name + '.png')
    canvas = Image.new('RGB', (1700, 820 if name == 'architecture' else 520), 'white')
    draw = ImageDraw.Draw(canvas)
    ft = ImageFont.truetype('C:/Windows/Fonts/msyh.ttc', 30)
    small = ImageFont.truetype('C:/Windows/Fonts/msyh.ttc', 25)
    def box(x, y, w, h, title, subtitle=''):
        draw.rounded_rectangle((x, y, x+w, y+h), radius=10, fill='#EEF3F7', outline='#607A8A', width=3)
        for text, yy, ff in [(title, y+22, ft), (subtitle, y+67, small)]:
            if text: draw.text((x+w/2, yy), text, fill='#182B38', font=ff, anchor='mt')
    def arrow(points):
        draw.line(points, fill='#3E657C', width=4)
        x,y=points[-1]; px,py=points[-2]
        if x>px: a=[(x,y),(x-16,y-9),(x-16,y+9)]
        elif x<px: a=[(x,y),(x+16,y-9),(x+16,y+9)]
        elif y>py: a=[(x,y),(x-9,y-16),(x+9,y-16)]
        else: a=[(x,y),(x-9,y+16),(x+9,y+16)]
        draw.polygon(a,fill='#3E657C')
    if name == 'architecture':
        for x,title,sub in [(35,'静态场景与输入','工具和冻结声明'),(600,'域生命周期','current candidate retired'),(1165,'统一 Pipeline','有界队列和三工作槽')]: box(x,25,500,125,title,sub)
        arrow([(535,88),(600,88)]);arrow([(1100,88),(1165,88)])
        box(1165,230,500,125,'同一 RD 有序工作流','编辑 水步 显示与提交');arrow([(1415,150),(1415,230)])
        box(600,230,500,125,'完整完成态','generation state_seq actual_time');arrow([(1165,292),(1100,292)])
        box(35,230,500,125,'异步回执与账本','CPU 消费 真实租约回收');arrow([(600,292),(535,292)])
        for x,title,sub in [(35,'材质与只读预览','床面 水面 白沫 亮纹'),(600,'高度与公共 Probe','worker 碰撞 Jolt 道具'),(1165,'保存与候选恢复','原始双槽和来源身份')]:
            box(x,520,500,125,title,sub);arrow([(850,355),(850,440),(x+250,440),(x+250,520)])
        draw.text((850,720),'完成态发布和消费者收敛具有独立生命周期',fill='#182B38',font=ft,anchor='mt')
    elif name == 'slots':
        titles=['FREE','CPU FILLING','GPU QUEUED','GPU COMPLETE','CPU RECLAIM']
        subtitles=['可以取得新租约','编码并上传','已记录待回读','真实结果可消费','等待回收确认']
        for i,(title,sub) in enumerate(zip(titles,subtitles)):
            x=20+i*340;box(x,30,300,125,title,sub)
            if i<4:arrow([(x+300,92),(x+340,92)])
        arrow([(1530,155),(1530,230),(170,230),(170,155)])
        draw.text((300,330),'正常回收需要同租约确认',font=small,fill='#182B38',anchor='mt')
        draw.text((300,370),'且最后读者已经结束',font=small,fill='#182B38',anchor='mt')
        box(600,340,500,125,'UNKNOWN 或 FATAL','保留租约 等待真实终结证明')
        arrow([(850,155),(850,340)])
    else:
        for i,(title,sub) in enumerate([('GPU 快照','借用 dirty 最多 16 块'),('异步回读','身份 有限值 版本校验'),('worker 准备','只回收真实完成任务'),('分帧 applying','最多 8 次 1500 μs 软预算')]):
            x=30+i*425;box(x,30,365,130,title,sub)
            if i<3:arrow([(x+365,95),(x+425,95)])
        box(1130,320,530,125,'全部消费或关闭丢弃','GPU ACK 后释放 borrowed')
        arrow([(1487,160),(1487,320)])
        arrow([(1130,382),(212,382),(212,160)])
        draw.text((690,410),'部分应用保留游标和真实读者',font=small,fill='#182B38',anchor='mt')
    canvas.save(path)
    return path

def get_abi():
    source = (ROOT / (BASE+'layout.gd')).read_text(encoding='utf-8-sig')
    tables = []
    for label, prefix in [('输入批次头', 'HEADER'), ('逐命令输入', 'COMMAND'), ('完成批次头', 'COMPLETION'), ('逐命令结果', 'RESULT')]:
        rows=[]
        for kind,size in [('U32',4),('I64',8),('F32',4),('HASH',32)]:
            m=re.search(r'^const '+prefix+'_'+kind+r'=(\{[^\n]+\})', source, re.M)
            if m:
                for field,offset in json.loads(m.group(1)).items(): rows.append([int(offset),size,kind,field])
        tables.append((label, sorted(rows)))
    return tables

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check',action='store_true')
    args=parser.parse_args()
    if args.check:
        index=json.loads(INDEX.read_text(encoding='utf-8'))
        drift=[p for p,h in index['sources'].items() if not (ROOT/p).is_file() or sha(ROOT/p)!=h]
        assert not drift, drift
        assert OUT.is_file() and sha(OUT)==index['document_sha256'], 'Document drift'
        print('TECHNICAL_PANORAMA_CHECK PASS');return
    WORK.mkdir(parents=True,exist_ok=True)
    doc=Document(); section=doc.sections[0]
    section.page_width=Mm(215.9);section.page_height=Mm(279.4)
    section.top_margin=Mm(19);section.bottom_margin=Mm(18)
    section.left_margin=Mm(20);section.right_margin=Mm(20)
    section.header_distance=Mm(8);section.footer_distance=Mm(9)
    for name,size in [('Normal',11),('Title',27),('Subtitle',12),('Heading 1',17),('Heading 2',13),('Caption',9)]:
        st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=BLACK
        st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
        pf=st.paragraph_format;pf.space_after=Pt(7);pf.line_spacing=1.2
        if name.startswith('Heading'):pf.space_before=Pt(15);pf.keep_with_next=True
        if name=='Normal':pf.widow_control=True
    # Remove inherited borders, color and rules from the title family.
    for name in ['Title','Subtitle']:
        ppr=doc.styles[name]._element.find(qn('w:pPr'))
        if ppr is not None:
            for item in list(ppr):
                if item.tag==qn('w:pBdr'):ppr.remove(item)
    header=section.header.paragraphs[0]
    font(header.add_run('SANDCASTLE   项目技术全景'),8)
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    font(footer.add_run('Sandcastle    '),8)
    fld=xml('fldSimple',instr='PAGE');footer._p.append(fld)
    doc.core_properties.title='Sandcastle 项目技术全景'
    doc.core_properties.subject='当前实现 架构 算法 数据 接口 生命周期与边界'
    doc.core_properties.author='Sandcastle'
    doc.core_properties.keywords='Godot RenderingDevice 沙水模拟 统一编辑'
    paragraph(doc,'Sandcastle 项目技术全景','Title')
    paragraph(doc,'架构 算法 数据接口 生命周期与运行边界','Subtitle')
    paragraph(doc,'内容依据当前源码和生效配置')
    paragraph(doc,'这份手册从场景入口到 GPU 权威状态，再到三维表现、碰撞、道具和原生检查点，解释当前项目如何运行。核心设计是让编辑、水 tick 和显示工作沿同一有序流推进，以完整完成态连接消费者，并用明确的所有权证明管理异步资源。')
    paragraph(doc,'内容依据项目当前源码、场景和实际被入口读取的配置。默认行为与可选能力分别说明；文档不包含开发流水、任务进度、性能测试流水或验收记录。阅读时可用 Word 标题导航或下页目录跳转，源文件定位见附录。')
    table(doc,['主题','读者能获得的内容'],[['架构与运行','模块职责、输入路径、CPU/GPU 边界与场景差异'],['数值与交互','浅水、悬沙、材料、编辑水响应和铲子账本'],['消费者与持久化','完整态绑定、碰撞分帧、Jolt 道具、保存恢复'],['工程接口','字节 ABI、公共入口、有效参数和来源文件']], [35,135])
    heading(doc,'内容导航',1).paragraph_format.page_break_before=True
    appendix=['附录 A 字节协议字段','附录 B 有效场景设置','附录 C main 材质动态输入','附录 D 公共接口索引','附录 E 技术来源索引']
    navigation=[f'{i:02d}  {s["title"]}' for i,s in enumerate(SECTIONS,1)]+appendix
    for i,title in enumerate(navigation,1):
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.05
        link=xml('hyperlink',anchor=f'section_{i}')
        r=xml('r');rpr=xml('rPr');rpr.append(xml('color',val='000000'));r.append(rpr)
        t=xml('t');t.text=title;r.append(t);link.append(r);p._p.append(link)
    source_list=list(dict.fromkeys(p for s in SECTIONS for p in s['sources']))
    def add_source(path):
        if path not in source_list:source_list.append(path)
        return f'S{source_list.index(path)+1:03d}'
    for i,s in enumerate(SECTIONS,1):
        section_heading=heading(doc,f'{i:02d}  {s["title"]}',1,(i,f'section_{i}'))
        if i==1:section_heading.paragraph_format.page_break_before=True
        for block in s['blocks']:
            if block['kind']=='p':paragraph(doc,block['text'])
            elif block['kind']=='h2':heading(doc,block['text'],2)
            elif block['kind']=='table':table(doc,block['headers'],block['rows'],block.get('widths'))
            elif block['kind']=='figure':
                p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
                run=p.add_run();run.add_picture(str(diagram(block['key'])),width=Mm(168))
                for prop in run._element.xpath('.//wp:docPr'):prop.set('descr',block['caption'])
                p=paragraph(doc,block['caption'],'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p=paragraph(doc,'源码与主题依据  '+'  '.join(add_source(x) for x in s['sources']))
        for run in p.runs:font(run,8.2)
    def appendix_heading(number,title,new_page=True):
        appendix_title=heading(doc,title,1,(number,f'section_{number}'))
        if new_page:appendix_title.paragraph_format.page_break_before=True
    number=len(SECTIONS)
    appendix_heading(number+1,appendix[0])
    paragraph(doc,'以下偏移直接读取 layout.gd，单位为字节，little-endian。U32 和 F32 为 4 B，I64 为 8 B 并以低高两个 uint32 传送，HASH 为 32 B。保留区不在表中列作可自由使用字段；外部调用应使用公共 Intent 与编码器，不能自行绕过语义校验。')
    add_source(BASE+'layout.gd')
    for label,rows in get_abi():
        heading(doc,label,2);table(doc,['偏移 B','长度 B','类型','字段'],rows,[20,20,25,105])
    appendix_heading(number+2,appendix[1])
    paragraph(doc,'下表来自各正式 manifest 的 metadata.settings。只列入口实际配置相关项，不把烘焙时的 metadata.configuration 或历史 checks 当作当前后端开关。数值求解器为空导出值时由 water_numerics.gd 解析为默认 MUSCL 二阶；历史检查点缺失数值字段仍按一阶解释。')
    keys=['water_policy','signed_speed_m_s','dig_amount_m3','fill_amount_m3','enable_erosion','enable_brush_water','foam_enabled','gravel_enabled','gravel_count','local_front_history','prototype_local_front_wide_field','render_subdivision','sand_cohesion_yield_height_m','use_player_camera']
    cases=['main','test1_geometry'];settings=[]
    for case in cases:
        path=f'assets/unified_edit/presentation/{case}.json';add_source(path)
        settings.append(json.loads((ROOT/path).read_text(encoding='utf-8'))['metadata']['settings'])
    rows=[[key]+[str(d.get(key,'未声明')) for d in settings] for key in keys]
    table(doc,['设置','main','geometry'],rows,[90,40,40])
    appendix_heading(number+3,appendix[2])
    paragraph(doc,'动态字段映射来自 main manifest，并由 material_plan 对照静态场景校验。每项列出 shader 和采样器到公共完成态字段的关系；材质静态参数以实际 .tscn 资源为准，不把旧烘焙参数快照视为运行中的全部有效值。')
    metadata=json.loads((ROOT/'assets/unified_edit/presentation/main.json').read_text(encoding='utf-8'))['metadata']
    grouped_materials={}
    for material in metadata['materials']:
        if not material['fields']:continue
        key=(material['shader'],tuple(sorted(material['fields'].items())))
        if key not in grouped_materials:grouped_materials[key]={'material':material,'ids':[]}
        grouped_materials[key]['ids'].append(material['id'])
    for group in grouped_materials.values():
        material=group['material']
        fields=material['fields']
        path=material['shader'].removeprefix('res://');sid=add_source(path)
        labels=' '.join(str(value) for value in group['ids'])
        heading(doc,f'材质 {labels}  {Path(path).stem}',2)
        paragraph(doc,f'{sid}  '+breakable(path)).paragraph_format.keep_with_next=True
        table(doc,['材质采样器','公共字段'],list(fields.items()),[105,65])
    appendix_heading(number+4,appendix[3],new_page=False)
    paragraph(doc,'接口签名从当前脚本提取。这里列出调用层和生命周期入口，不鼓励业务层直接访问内部 RD 或数值录制方法。调用返回字典或错误名后，仍需按接纳与终结结果的区别处理。')
    api={
      'scripts/beach/unified_playable_entry.gd':['save_gpu_checkpoint','restore_gpu_checkpoint','shutdown'],
      BASE+'player_input.gd':['begin','move','finish'],
      BASE+'domain_runtime.gd':['boot','submit_intent','submit_control','submit_display_intent','result','take_result','request_tick','request_wave','request_preview','request_probe','save_checkpoint','restore_checkpoint','cancel_restore','shutdown'],
      'scripts/props/floating_props_service.gd':['attach','manipulate_body','set_paused','reset','shutdown'],
      BASE+'configuration.gd':['freeze','validate'],
    }
    for path,names in api.items():
        sid=add_source(path);heading(doc,Path(path).stem,2)
        paragraph(doc,f'{sid}  '+breakable(path)).paragraph_format.keep_with_next=True
        text=(ROOT/path).read_text(encoding='utf-8-sig');rows=[]
        for match in re.finditer(r'^(?:static )?func (\w+)\(([^\n]*)',text,re.M):
            if match.group(1) in names:rows.append([match.group(1),match.group(0).removeprefix('static ').removeprefix('func ').rstrip(':')])
        assert {row[0] for row in rows} == set(names), f'Public API drift: {path}'
        table(doc,['方法','当前签名'],rows,[38,132])
    appendix_heading(number+5,appendix[4])
    paragraph(doc,'编号用于正文定位。路径相对项目根目录；逐文件 SHA256 另存 technical-panorama.sources.json，可检查本文生成所依据的来源是否变化。来源索引不包含历史验收状态，文档中的技术说明以对应源码和实际配置为依据。')
    table(doc,['编号','当前源文件'],[[f'S{i:03d}',p] for i,p in enumerate(source_list,1)],[18,152])
    tracked=source_list+['tools/docs/build_technical_panorama.py','tools/docs/technical_panorama_content.py']
    missing=[p for p in tracked if not (ROOT/p).is_file()]
    assert not missing, missing
    OUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUT)
    index={'schema':1,'document':OUT.relative_to(ROOT).as_posix(),'source_scope':'Current implementation, configuration, interfaces and technical documents only','sources':{p:sha(ROOT/p) for p in sorted(set(tracked))},'document_sha256':sha(OUT),'sections':len(SECTIONS),'appendices':len(appendix)}
    INDEX.write_text(json.dumps(index,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'document':str(OUT),'sections':len(SECTIONS),'appendices':len(appendix),'sources':len(index['sources'])},ensure_ascii=False))

if __name__=='__main__':main()
